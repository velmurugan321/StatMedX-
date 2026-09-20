"""Module 14 — Result export: CSV, Excel, Word (docx), printable HTML (→PDF via browser)."""
from __future__ import annotations

import io

import pandas as pd


def _tables(result: dict) -> list[tuple[str, pd.DataFrame]]:
    out = []
    seen = {}
    for b in result.get("blocks", []):
        if b.get("type") == "table":
            name = b.get("name", "table")
            if name in seen:
                seen[name] += 1
                name = f"{name} {seen[name]}"
            else:
                seen[name] = 0
            df = pd.DataFrame(b.get("rows", []), columns=[c["label"] for c in b.get("columns", [])])
            out.append((name[:31] or "table", df))
    return out


def to_csv(result: dict) -> bytes:
    parts = []
    for name, df in _tables(result):
        parts.append(f"# {name}\n")
        parts.append(df.to_csv(index=False))
    return "\n".join(parts).encode()


def to_excel(result: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        first = True
        for name, df in _tables(result):
            df.to_excel(w, sheet_name=name or "results", index=False,
                        startrow=1 if not first else 0)
            first = False
        if first:
            pd.DataFrame({"StatMedX": ["no tables in this result"]}).to_excel(w, index=False)
    return buf.getvalue()


def to_docx(result: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(result.get("title", "StatMedX result"), level=2)
    if result.get("command"):
        p = doc.add_paragraph()
        r = p.add_run(f"Command: {result['command']}")
        r.font.name = "Courier New"
        r.font.size = Pt(9)
    for b in result.get("blocks", []):
        t = b.get("type")
        if t == "text":
            # strip simple html tags
            import re
            txt = re.sub(r"<[^>]+>", "", b.get("content", ""))
            doc.add_paragraph(txt)
        elif t == "table":
            name, df = b.get("name", ""), pd.DataFrame(
                b.get("rows", []), columns=[c["label"] for c in b.get("columns", [])])
            if name:
                doc.add_heading(name, level=4)
            table = doc.add_table(rows=1, cols=max(1, len(df.columns)))
            table.style = "Light Grid Accent 1"
            for j, c in enumerate(df.columns):
                table.rows[0].cells[j].text = str(c)
            for _, row in df.iterrows():
                cells = table.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = "" if v is None else str(v)
            if b.get("note"):
                doc.add_paragraph(b["note"]).runs and None
        elif t == "figure":
            doc.add_paragraph(f"[figure: {b.get('name', '')} — view in StatMedX or print to PDF]")
    if result.get("blocks") and not _tables(result):
        doc.add_paragraph("(text result)")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_html(result: dict) -> bytes:
    import html as _html
    import re

    parts = [f"""<!doctype html><html><head><meta charset="utf-8"><title>{_html.escape(result.get('title','StatMedX result'))}</title>
<style>
body{{font-family:Inter,system-ui,sans-serif;max-width:900px;margin:40px auto;padding:0 20px;color:#1e293b}}
h1{{font-size:20px}} .cmd{{font-family:monospace;background:#f1f5f9;padding:8px 12px;border-radius:6px;font-size:13px}}
table{{border-collapse:collapse;margin:12px 0;width:100%;font-size:13px}}
th,td{{border:1px solid #cbd5e1;padding:6px 10px;text-align:left}}
th{{background:#f1f5f9}} .note{{color:#64748b;font-size:12px;margin-top:-8px}}
h3{{margin:18px 0 4px;font-size:15px}} @media print{{body{{margin:10px auto}}}}
</style></head><body>"""]
    parts.append(f"<h1>{_html.escape(result.get('title', ''))}</h1>")
    if result.get("command"):
        parts.append(f"<div class='cmd'>{_html.escape(result['command'])}</div>")
    for b in result.get("blocks", []):
        t = b.get("type")
        if t == "text":
            parts.append(re.sub(r"<(?!/?(b|i|br|code|sub|sup)\b)[^>]+>", "", b.get("content", "")))
        elif t == "table":
            parts.append(f"<h3>{_html.escape(b.get('name',''))}</h3><table><tr>")
            for c in b.get("columns", []):
                parts.append(f"<th>{_html.escape(c['label'])}</th>")
            parts.append("</tr>")
            for row in b.get("rows", []):
                parts.append("<tr>" + "".join(f"<td>{_html.escape('' if v is None else str(v))}</td>" for v in row) + "</tr>")
            parts.append("</table>")
            if b.get("note"):
                parts.append(f"<div class='note'>{_html.escape(b['note'])}</div>")
        elif t == "figure":
            parts.append(f"<p><i>[chart: {_html.escape(b.get('name',''))} — charts are embedded in the app; "
                         f"use screenshot or the app's print view for figures]</i></p>")
    parts.append("<p style='margin-top:30px;color:#94a3b8;font-size:11px'>Generated by StatMedX — use Print → Save as PDF</p>")
    parts.append("</body></html>")
    return "\n".join(parts).encode()


def export(result: dict, fmt: str) -> tuple[bytes, str, str]:
    fmt = fmt.lower()
    if fmt == "csv":
        return to_csv(result), "text/csv", "csv"
    if fmt in ("excel", "xlsx"):
        return to_excel(result), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    if fmt in ("word", "docx"):
        return to_docx(result), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    if fmt in ("pdf", "html"):
        return to_html(result), "text/html", "html"
    raise ValueError(f"Unknown export format '{fmt}'")
